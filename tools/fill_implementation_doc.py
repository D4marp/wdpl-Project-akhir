from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SRC = Path(
    r"C:\Users\WINDOWS\Downloads\Kelompok 5 - Penerapan Factory Pattern dan Observer Pattern pada Aplikasi Pencatatan Keuangan UMKM berbasis Web dan Mobile menggunakan Go (Gin Framework), Flutter, dan Next.js.docx"
)
OUT = Path(
    r"D:\LARAGON\laragon\www\wdpl-Project-akhir\Kelompok 5 - Implementasi Factory Observer UMKM - revisi.docx"
)


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def add_code_after(anchor, code):
    p = insert_paragraph_after(anchor, "")
    p.style = "Normal"
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(code)
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p


def add_bullets_after(anchor, items):
    cur = anchor
    for item in items:
        cur = insert_paragraph_after(cur, item, "List Bullet")
    return cur


def add_numbered_after(anchor, items):
    cur = anchor
    for item in items:
        cur = insert_paragraph_after(cur, item, "List Number")
    return cur


def add_table_after(anchor, rows):
    table = anchor._parent.add_table(rows=1, cols=len(rows[0]), width=Inches(6.0))
    table._tbl.getparent().remove(table._tbl)
    anchor._p.addnext(table._tbl)

    hdr = table.rows[0].cells
    for i, value in enumerate(rows[0]):
        hdr[i].text = value
        for run in hdr[i].paragraphs[0].runs:
            run.bold = True

    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    return table


def style_document(doc):
    for style_name in ["Normal", "Body Text"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Arial"
            style.font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Arial"


def main():
    doc = Document(SRC)
    style_document(doc)

    paragraphs = doc.paragraphs
    start_idx = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "Bab 12 : Implementasi")

    # Remove the placeholder heading and unrelated stray paragraph before bibliography.
    for idx in sorted([start_idx + 1, start_idx + 2], reverse=True):
        if idx < len(doc.paragraphs):
            delete_paragraph(doc.paragraphs[idx])

    anchor = doc.paragraphs[start_idx]

    intro = (
        "Bab ini menjelaskan implementasi aplikasi pencatatan keuangan UMKM yang terdapat pada folder "
        "wdpl-Project-akhir. Implementasi dibuat dalam tiga bagian utama, yaitu backend berbasis Go dengan "
        "Gin Framework, web client berbasis Next.js, dan mobile client berbasis Flutter. Fokus utama bab ini "
        "adalah menunjukkan bagaimana Factory Pattern dan Observer Pattern diterapkan langsung dalam kode "
        "aplikasi, bukan hanya sebagai konsep teoritis."
    )
    anchor = insert_paragraph_after(anchor, intro)

    sections = [
        (
            "A. Struktur Folder Implementasi",
            [
                "Folder backend berisi REST API Go. Bagian ini menangani koneksi database, middleware tenant, handler HTTP, service transaksi, model GORM, serta implementasi Factory Pattern dan Observer Pattern.",
                "Folder web berisi dashboard Next.js. Bagian ini menyediakan halaman ringkasan, transaksi, laporan, dan notifikasi yang mengambil data dari REST API backend.",
                "Folder mobile berisi aplikasi Flutter. Bagian ini memakai Provider untuk state management dan ApiService untuk komunikasi HTTP ke backend.",
            ],
            None,
        ),
        (
            "B. Implementasi Backend Go Gin",
            [
                "Backend menggunakan Gin sebagai router HTTP, GORM sebagai ORM, dan MySQL sebagai database. File utama berada pada backend/cmd/main.go. Pada saat aplikasi dijalankan, program memuat konfigurasi dari .env, membuka koneksi database, menjalankan auto migration, membuat instance factory, membuat event bus observer, lalu mendaftarkan endpoint API.",
                "Endpoint API memakai header X-Tenant-ID untuk memisahkan data antar UMKM. Middleware tenant membaca header tersebut dan menyimpannya ke context Gin sehingga handler tidak perlu membaca header secara berulang.",
            ],
            """r := gin.Default()
r.Use(corsMiddleware())
api := r.Group("/api")
api.Use(middleware.TenantMiddleware())
api.POST("/transactions", txHandler.Create)
api.GET("/transactions", txHandler.List)
api.GET("/transactions/summary", txHandler.Summary)
api.GET("/reports", reportHandler.Get)
api.POST("/budgets", budgetHandler.Create)
api.GET("/budgets", budgetHandler.List)
api.GET("/notifications", notifHandler.List)""",
        ),
        (
            "C. Implementasi Model dan Database",
            [
                "Model utama terdiri atas Tenant, Transaction, Budget, dan Notification. Auto migration dijalankan pada saat backend start sehingga tabel dibuat otomatis jika belum tersedia.",
                "Tenant digunakan sebagai identitas UMKM. Transaction menyimpan data pemasukan dan pengeluaran. Budget menyimpan batas anggaran per kategori. Notification menyimpan informasi transaksi dan peringatan budget.",
            ],
            """db.AutoMigrate(
    &models.Tenant{},
    &models.Transaction{},
    &models.Budget{},
    &models.Notification{},
)""",
        ),
        (
            "D. Implementasi Factory Pattern pada Transaksi",
            [
                "Factory Pattern diterapkan pada file backend/internal/patterns/factory.go. Interface Transaction menjadi kontrak untuk semua jenis transaksi. IncomeTransaction dan ExpenseTransaction menjadi concrete product. TransactionFactory bertugas membuat objek transaksi berdasarkan nilai type yang dikirim client.",
                "Handler tidak membuat IncomeTransaction atau ExpenseTransaction secara langsung. Handler hanya meneruskan request ke TransactionService, kemudian service meminta TransactionFactory membuat objek yang sesuai. Dengan cara ini, logika pembuatan objek transaksi berada pada satu tempat dan lebih mudah diperluas.",
            ],
            """type Transaction interface {
    GetType() string
    GetCategory() string
    GetDescription() string
    GetAmount() float64
    GetTenantID() string
    Validate() error
}

type TransactionFactory struct{}

func (f *TransactionFactory) Create(
    txType, category, description, tenantID string,
    amount float64,
) (Transaction, error) {
    switch txType {
    case "income":
        return &IncomeTransaction{Amount: amount, Category: category, Description: description, TenantID: tenantID}, nil
    case "expense":
        return &ExpenseTransaction{Amount: amount, Category: category, Description: description, TenantID: tenantID}, nil
    default:
        return nil, fmt.Errorf("tipe transaksi '%s' tidak dikenal", txType)
    }
}""",
        ),
        (
            "E. Implementasi Factory Pattern pada Laporan",
            [
                "Selain transaksi, Factory Pattern juga digunakan untuk laporan. ReportFactory membuat objek DailyReport, WeeklyReport, atau MonthlyReport berdasarkan parameter period dari endpoint /api/reports.",
                "Setiap report memiliki metode Generate yang sama, tetapi rentang waktunya berbeda. DailyReport mengambil transaksi satu hari terakhir, WeeklyReport tujuh hari terakhir, dan MonthlyReport satu bulan terakhir.",
            ],
            """type Report interface {
    Generate(tenantID string, db *gorm.DB) (ReportData, error)
    GetPeriod() string
}

func (f *ReportFactory) Create(period string) (Report, error) {
    switch period {
    case "daily":
        return &DailyReport{}, nil
    case "weekly":
        return &WeeklyReport{}, nil
    case "monthly":
        return &MonthlyReport{}, nil
    default:
        return nil, fmt.Errorf("periode tidak valid")
    }
}""",
        ),
        (
            "F. Implementasi Observer Pattern",
            [
                "Observer Pattern diterapkan pada backend/internal/patterns/observer.go. Subject atau publisher pada implementasi ini adalah TransactionEventBus. Observer yang didaftarkan adalah BudgetObserver, NotificationObserver, dan AuditObserver.",
                "Ketika transaksi berhasil disimpan, TransactionService memanggil eventBus.Publish. Event tersebut kemudian dikirim ke seluruh observer secara asynchronous menggunakan goroutine. Dengan pendekatan ini, response API tidak perlu menunggu proses notifikasi atau audit selesai.",
            ],
            """func (eb *TransactionEventBus) Publish(event TransactionEvent) {
    eb.mu.RLock()
    observers := make([]TransactionObserver, len(eb.observers))
    copy(observers, eb.observers)
    eb.mu.RUnlock()

    for _, obs := range observers {
        obs := obs
        go func() {
            defer func() {
                if r := recover(); r != nil {
                    fmt.Printf("[EventBus] panic di %s: %v\\n", obs.Name(), r)
                }
            }()
            obs.OnTransaction(event)
        }()
    }
}""",
        ),
        (
            "G. Alur Create Transaction",
            [
                "Alur utama aplikasi terjadi saat pengguna menambahkan transaksi. Client mengirim request POST /api/transactions ke backend. Handler membaca tenant dari middleware dan meneruskan payload ke TransactionService.",
                "TransactionService menggunakan TransactionFactory untuk membuat objek transaksi, menjalankan validasi bisnis, menyimpan record ke database, lalu mempublish TransactionEvent ke EventBus. Setelah event dipublish, BudgetObserver memperbarui spent_amount untuk kategori terkait, NotificationObserver membuat notifikasi transaksi baru, dan AuditObserver menulis log audit.",
            ],
            """record := &models.Transaction{
    TenantID:    tx.GetTenantID(),
    Type:        tx.GetType(),
    Category:    tx.GetCategory(),
    Description: tx.GetDescription(),
    Amount:      tx.GetAmount(),
    CreatedAt:   time.Now(),
}
s.db.Create(record)
s.eventBus.Publish(patterns.TransactionEvent{
    TenantID: record.TenantID,
    Type: record.Type,
    Amount: record.Amount,
    Category: record.Category,
    Description: record.Description,
    CreatedAt: record.CreatedAt,
})""",
        ),
        (
            "H. Implementasi Web Client Next.js",
            [
                "Web client berada pada folder web. File web/lib/api.ts menjadi pusat komunikasi API. Nilai NEXT_PUBLIC_API_URL diambil dari .env.local dan default-nya adalah http://localhost:8080. Header X-Tenant-ID selalu dikirim agar backend dapat memisahkan data tenant.",
                "Halaman utama menampilkan ringkasan saldo, pemasukan, pengeluaran, dan akses ke halaman transaksi. Halaman transaksi menyediakan form pencatatan dan riwayat transaksi. Halaman laporan mengambil data dari endpoint /api/reports dan menampilkannya dalam bentuk visualisasi. Halaman notifikasi mengambil data dari endpoint /api/notifications.",
            ],
            """const BASE = process.env.NEXT_PUBLIC_API_URL ?? "http://localhost:8080";
const TENANT = process.env.NEXT_PUBLIC_TENANT_ID ?? "tenant-001";

const h = { "Content-Type": "application/json", "X-Tenant-ID": TENANT };

async function req<T>(path: string, init?: RequestInit): Promise<T> {
    const res = await fetch(`${BASE}${path}`, { ...init, headers: h });
    return res.json();
}""",
        ),
        (
            "I. Implementasi Mobile Client Flutter",
            [
                "Mobile client berada pada folder mobile. File lib/config.dart menyimpan konfigurasi API melalui dart-define. Untuk Android emulator, API_URL menggunakan http://10.0.2.2:8080 karena localhost pada emulator mengarah ke emulator itu sendiri, bukan komputer host.",
                "State aplikasi dikelola oleh FinanceProvider. Provider memanggil ApiService untuk mengambil summary, transaksi, laporan, dan notifikasi. Dengan pola ini, layar Flutter tidak berkomunikasi langsung dengan HTTP client, tetapi melalui provider sehingga state lebih terpusat.",
            ],
            """class AppConfig {
  static const String apiUrl = String.fromEnvironment(
    'API_URL',
    defaultValue: 'http://10.0.2.2:8080',
  );
  static const String tenantId = String.fromEnvironment(
    'TENANT_ID',
    defaultValue: 'tenant-001',
  );
}""",
        ),
        (
            "J. Integrasi Multi-Platform",
            [
                "Backend menjadi pusat data dan aturan bisnis. Web client dan mobile client tidak menyimpan logika Factory Pattern maupun Observer Pattern secara langsung. Keduanya hanya mengirim request ke REST API dan menampilkan response.",
                "Pendekatan ini membuat implementasi pattern tetap terpusat di backend. Jika aturan validasi transaksi atau mekanisme observer berubah, web dan mobile tidak perlu banyak diubah selama kontrak API tetap sama.",
            ],
            None,
        ),
        (
            "K. Cara Menjalankan Implementasi",
            [
                "Backend dijalankan dari folder backend. File .env perlu disesuaikan dengan konfigurasi MySQL lokal. Setelah database umkm_finance dibuat, backend dapat dibuild dan dijalankan melalui backend.exe.",
                "Web dijalankan dari folder web menggunakan npm install dan npm run dev. Aplikasi dapat dibuka melalui http://localhost:3000.",
                "Mobile dijalankan dari folder mobile menggunakan flutter pub get dan flutter run. Untuk Android emulator gunakan API_URL=http://10.0.2.2:8080, sedangkan untuk Chrome gunakan API_URL=http://localhost:8080.",
            ],
            """Backend:
go build -o ./bin/backend.exe ./cmd/main.go
./bin/backend.exe

Web:
npm install
npm run dev

Mobile Android Emulator:
flutter run --dart-define=API_URL=http://10.0.2.2:8080 --dart-define=TENANT_ID=tenant-001

Mobile Web/Chrome:
flutter run -d chrome --dart-define=API_URL=http://localhost:8080 --dart-define=TENANT_ID=tenant-001""",
        ),
        (
            "L. Ringkasan Hasil Implementasi",
            [
                "Factory Pattern berhasil diterapkan untuk memisahkan pembuatan objek transaksi dan laporan dari handler HTTP.",
                "Observer Pattern berhasil diterapkan untuk memisahkan proses lanjutan setelah transaksi, yaitu pembaruan budget, pembuatan notifikasi, dan audit log.",
                "Backend, web, dan mobile terhubung melalui REST API dengan header X-Tenant-ID sebagai mekanisme multi-tenant sederhana.",
                "Struktur project mendukung maintainability karena setiap layer memiliki tanggung jawab yang jelas: handler menerima request, service menjalankan proses bisnis, pattern mengatur variasi objek dan event, sedangkan client hanya berfokus pada tampilan dan interaksi pengguna.",
            ],
            None,
        ),
    ]

    for title, paras, code in sections:
        anchor = insert_paragraph_after(anchor, title, "Heading 2")
        for para in paras:
            anchor = insert_paragraph_after(anchor, para)
        if code:
            anchor = add_code_after(anchor, code)

    # Add compact evidence table after section L.
    anchor = insert_paragraph_after(anchor, "M. Pemetaan File Implementasi", "Heading 2")
    anchor = insert_paragraph_after(
        anchor,
        "Tabel berikut merangkum file utama yang menjadi bukti implementasi pada project wdpl-Project-akhir.",
    )
    table = add_table_after(
        anchor,
        [
            ["Bagian", "File", "Fungsi"],
            ["Factory Pattern", "backend/internal/patterns/factory.go", "Membuat transaksi income/expense dan laporan daily/weekly/monthly."],
            ["Observer Pattern", "backend/internal/patterns/observer.go", "Event bus dan observer untuk budget, notifikasi, dan audit."],
            ["Service Transaksi", "backend/internal/services/transaction_service.go", "Menghubungkan factory, database, dan event bus."],
            ["REST API", "backend/cmd/main.go", "Registrasi router, middleware tenant, CORS, dan handler API."],
            ["Web Client", "web/lib/api.ts", "Helper fetch API dengan X-Tenant-ID."],
            ["Mobile Client", "mobile/lib/services/api_service.dart", "Helper HTTP Flutter untuk summary, transaksi, laporan, dan notifikasi."],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
